"""Writing rows into an Excel file, without repeating what is already there."""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

DATE_COLUMN = "collected_at"


def _header(columns: list[str]) -> list[str]:
    return [*columns, DATE_COLUMN]


def existing_keys(path: Path, sheet: str, key_field: str) -> set[str]:
    """Read the keys already saved, so later runs don't duplicate rows."""
    if not path.exists():
        return set()
    workbook = load_workbook(path, read_only=True)
    if sheet not in workbook.sheetnames:
        workbook.close()
        return set()
    worksheet = workbook[sheet]
    rows = worksheet.iter_rows(values_only=True)
    try:
        header = next(rows)
    except StopIteration:
        workbook.close()
        return set()
    if key_field not in header:
        workbook.close()
        return set()
    index = header.index(key_field)
    keys = {
        str(row[index]).strip() for row in rows if row and row[index] is not None
    }
    workbook.close()
    return keys


def save(
    path: Path,
    sheet: str,
    columns: list[str],
    items: list[dict[str, str]],
) -> int:
    """Append the items to the spreadsheet and return how many rows were added."""
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.exists():
        workbook = load_workbook(path)
        worksheet = (
            workbook[sheet] if sheet in workbook.sheetnames else workbook.create_sheet(sheet)
        )
        if (
            worksheet.max_row == 1
            and worksheet.max_column == 1
            and worksheet["A1"].value is None
        ):
            worksheet.append(_header(columns))
    else:
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = sheet
        worksheet.append(_header(columns))

    for cell in worksheet[1]:
        cell.font = Font(bold=True)
    worksheet.freeze_panes = "A2"

    now = datetime.now().strftime("%Y-%m-%d %H:%M")
    for item in items:
        worksheet.append([item.get(column, "") for column in columns] + [now])

    # Column width from the longest value, capped so it stays readable.
    for index, column in enumerate(_header(columns), start=1):
        widest = len(str(column))
        for item in items:
            value = item.get(column, "") if column != DATE_COLUMN else now
            widest = max(widest, len(str(value)))
        worksheet.column_dimensions[get_column_letter(index)].width = min(widest + 2, 60)

    workbook.save(path)
    return len(items)
